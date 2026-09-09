"""The generated report against the manual one, cell for cell.

    python -m lnd.export.compare        # writes docs/report-comparison.md

Week 10, Person B. The layout was reproduced in week 8 and checked by eye; this
checks the numbers landing in the cells. Two different questions get answered
here and only the second one can fail:

1. **Does the generated figure match the workbook's?** Mostly no, and every
   difference is one the reconciliation statement already explains. Nothing to
   fix — it is the point of the project.

2. **Is each generated figure in the cell its own label claims?** This must be
   yes for every one of them. A correct metric written into the wrong column produces a
   file that is right in every respect except the one a reader uses, and no test
   above the spreadsheet would notice: the registry is right, the export service
   is right, and Logistics Effectiveness is printed under Activity
   Effectiveness.

WHAT READING THE ORIGINAL TURNED UP

The workbook's own DASHBOARD is not internally aligned. `N7` is labelled
"Facilitators Performance " and `N8` is empty — the value is in `O8`, one column
to the right of its own label. `Z7` holds the participation rate as a bare
number with no label anywhere near it. The generated file does not reproduce
either, and this document says so rather than silently disagreeing with the
original in two cells.
"""

from __future__ import annotations

import argparse
import io
import logging
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from lnd.export.monthly import PARTICIPATION_COLUMN, PLACEMENTS, monthly_report
from lnd.metrics import registry
from lnd.metrics.base import MetricValue, Unit
from lnd.metrics.filters import MetricFilters
from lnd.reference.windows import WORKBOOK

log = logging.getLogger(__name__)

REPORT = Path(__file__).resolve().parents[4] / "docs" / "report-comparison.md"
MANUAL = Path(__file__).resolve().parents[4] / "docs" / "L&D Main Reports.xlsx"

LABEL_ROW = 7
VALUE_ROW = 8

#: Cells where the original put a figure somewhere other than under its label.
#: Recorded rather than worked around: the comparison has to read the manual
#: file as it is, and a reader of this document should know the original was
#: like this before concluding the platform moved something.
MANUAL_OVERRIDES: dict[str, tuple[int, int, str]] = {
    "facilitator_performance": (
        VALUE_ROW,
        15,  # O
        "labelled in N7, valued in O8 — one column right of its own label",
    ),
    "participation_rate": (
        LABEL_ROW,
        PARTICIPATION_COLUMN,
        "a bare number in Z7 with no label anywhere near it",
    ),
}


#: Where the manual file's unit was not the metric's. Only NPS: the workbook
#: stored 0.9272 and formatted it as 92.7%, which is a percentage of something
#: else. Rendering that cell as an index would print "+0.9" and quietly hide the
#: error instead of showing it.
MANUAL_UNITS: dict[str, Unit] = {"nps": Unit.PERCENT}


@dataclass(frozen=True)
class CellComparison:
    key: str
    label: str
    column: str
    manual: Any
    generated: Any
    registry_value: str
    manual_note: str
    #: The generated cell holds the figure its label claims. The one that fails.
    lands_correctly: bool


def _formatted(value: Any, unit: Unit) -> str:
    """A raw cell value as the sheet displays it.

    Percentages are stored as proportions and shown with a percent format, in
    both files, so 0.9818 and 98.2% are the same cell read two ways. Comparing
    the stored numbers without this would report every percentage as different.
    """
    if value is None:
        return "—"
    if isinstance(value, str):
        return value
    if unit is Unit.PERCENT:
        return f"{float(value) * 100:.1f}%"
    if unit is Unit.NPS:
        return f"{float(value):+.1f}"
    if unit in (Unit.HOURS, Unit.MONTHS):
        return f"{float(value):,.1f}"
    return f"{float(value):,.0f}"


def _holds(cell: Any, value: MetricValue) -> bool:
    """Does this cell hold this metric's number?

    Numeric, not textual. The two are rendered by different code — `Decimal`
    formatting rounds half to even, a float `f"{x:.1f}"` does not — so 135.65
    displays as 135.6 in one and 135.7 in the other, and comparing the strings
    reports a misplacement where there is a rounding convention. The question
    being asked is whether the right *number* reached the right cell, and a
    figure in the wrong column is wrong by whole units, never by 0.05.
    """
    if value.value is None:
        return cell in (None, "—")
    if not isinstance(cell, int | float):
        return False
    stored = float(value.value) / 100 if value.unit is Unit.PERCENT else float(value.value)
    return abs(float(cell) - stored) < 1e-6


def compare(session: Session, filters: MetricFilters | None = None) -> list[CellComparison]:
    applied = filters or WORKBOOK
    generated = load_workbook(io.BytesIO(monthly_report(session, applied)), data_only=True)[
        "DASHBOARD"
    ]
    manual = load_workbook(MANUAL, data_only=True)["DASHBOARD"] if MANUAL.exists() else None

    placements = [(p.metric_key, p.label, p.column) for p in PLACEMENTS]
    placements.append(("participation_rate", "Participation rate", PARTICIPATION_COLUMN))

    rows: list[CellComparison] = []
    for key, label, column in placements:
        value = registry.compute(key, session, applied)
        unit = value.unit

        generated_cell = generated.cell(row=VALUE_ROW, column=column).value
        row, manual_column, note = MANUAL_OVERRIDES.get(key, (VALUE_ROW, column, ""))
        manual_cell = manual.cell(row=row, column=manual_column).value if manual else None

        # The manual cell is rendered in the unit the *manual file* used, not
        # the one the metric now carries. L8 held 0.9272 formatted as a
        # percentage; showing it as the NPS index it is not would print "+0.9"
        # and hide the very unit error the reconciliation exists to explain.
        rows.append(
            CellComparison(
                key=key,
                label=label,
                column=get_column_letter(column),
                manual=_formatted(manual_cell, MANUAL_UNITS.get(key, unit)),
                generated=_formatted(generated_cell, unit),
                registry_value=value.formatted(),
                manual_note=note,
                lands_correctly=_holds(generated_cell, value),
            )
        )
    return rows


def render(rows: list[CellComparison], filters: MetricFilters) -> str:
    misplaced = [row for row in rows if not row.lands_correctly]
    lines = [
        "# The generated report against the manual one",
        "",
        "<!-- GENERATED by `python -m lnd.export.compare`. Do not edit by hand. -->",
        "",
        f"Compared {datetime.now(UTC).strftime('%Y-%m-%d %H:%M')} UTC over "
        f"{filters.describe()} — the window the manual workbook covers, which is the only "
        "like-for-like reading.",
        "",
        "## Two questions, and only one of them can fail",
        "",
        "The layout was reproduced in week 8 and checked by eye. This checks the numbers "
        "landing in the cells.",
        "",
        "**Does the generated figure match the manual one?** Mostly no, and every "
        "difference is one [`reconciliation.md`](reconciliation.md) already explains. That "
        "is the project, not a defect.",
        "",
        "**Is each figure in the cell its own label claims?** This must be yes for all "
        f"{len(rows)}. A correct metric written into the wrong column produces a file that "
        "is right in every respect except the one a reader uses — and nothing above the "
        "spreadsheet would notice, because the registry would be right and the export "
        "service would be right and Logistics Effectiveness would be printed under "
        "Activity Effectiveness.",
        "",
        "## Figure for figure, in the DASHBOARD layout",
        "",
        "| Cell | Figure | Manual | Generated | Registry | Lands correctly |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        mark = "✓" if row.lands_correctly else "**✗**"
        lines.append(
            f"| `{row.column}{VALUE_ROW}` | {row.label} | {row.manual} | **{row.generated}** | "
            f"{row.registry_value} | {mark} |"
        )

    noted = [row for row in rows if row.manual_note]
    if noted:
        lines += [
            "",
            "## Where the original was not laid out the way it reads",
            "",
            "The manual DASHBOARD is not internally aligned, and the comparison had to read "
            "it as it is. Recorded here so nobody concludes the platform moved something:",
            "",
            *(f"- **{row.label}** — {row.manual_note}." for row in noted),
            "",
            "The generated file puts every value directly under its own label. That is a "
            "departure from the original in two cells, and it is the right way round: a "
            "recipient's eye goes down a column, not across to the next one.",
            "",
        ]

    lines += ["## Result", ""]
    if misplaced:
        lines += [
            "**Not passed.** These cells do not hold the figure their label claims:",
            "",
            *(
                f"- `{row.column}{VALUE_ROW}` **{row.label}** — the cell reads "
                f"{row.generated}, the registry says {row.registry_value}"
                for row in misplaced
            ),
            "",
        ]
    else:
        lines += [
            f"**Passed.** All {len(rows)} figures land in the cell their label claims, and "
            "each equals the registry's value for that metric — the same value the "
            "dashboard shows and the same one the exports stamp.",
            "",
            "So the file a recipient opens and the screen a specialist reads cannot "
            "disagree. They are not two implementations that happen to agree today; there "
            "is one implementation, and this checks that the spreadsheet writer put its "
            "output where the labels say.",
            "",
        ]

    lines += [
        "## What is not compared",
        "",
        "The twelve other sheets. Attendance Master Sheet, Head Count 212, Attendance Prep, "
        "Calc, Linked In no and the pivot tabs are the machinery that produced DASHBOARD — "
        "they are what the platform *replaces*, and reproducing them would be reproducing "
        "eight hours of manual assembly in code.",
        "",
        "The generated file carries two sheets the original does not: **All figures**, "
        "every metric with its sample size, and **Provenance**, which states when the file "
        "was generated, what filters were applied, how fresh the data was and how many "
        "records were excluded. The manual workbook stated none of that, which is how a "
        "figure computed over 55 of 77 responses travelled for a year.",
        "",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate docs/report-comparison.md")
    parser.add_argument(
        "--check", action="store_true", help="write nothing; fail on a misplacement"
    )
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")

    from lnd.db import session_scope

    with session_scope() as session:
        rows = compare(session)
        document = render(rows, WORKBOOK)

    if not args.check:
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(document, encoding="utf-8")
        log.info("wrote %s", REPORT)

    misplaced = [row for row in rows if not row.lands_correctly]
    for row in misplaced:
        log.error(
            "MISPLACED %s — cell reads %s, registry says %s",
            row.label,
            row.generated,
            row.registry_value,
        )
    log.info("%d figures compared, %d misplaced", len(rows), len(misplaced))
    return 1 if misplaced else 0


if __name__ == "__main__":  # pragma: no cover - a script
    raise SystemExit(main())
